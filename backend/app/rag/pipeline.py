"""RAG pipeline: extract -> clean -> chunk -> embed -> store -> retrieve."""
from __future__ import annotations

import hashlib
import io
import re
from dataclasses import dataclass
from pathlib import Path

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.ai.embedder import cosine, get_embedder
from app.core.errors import AppError
from app.models.tables import DocumentChunk, Subject, SyllabusDocument
from app.storage.backend import get_storage, syllabus_key


class DocumentError(AppError):
    def __init__(self, message: str) -> None:
        super().__init__("DOCUMENT_ERROR", message, 400)


@dataclass
class ExtractedDocument:
    text: str
    pages: int


def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[\u0000-\u0008\u000b\u000c\u000e-\u001f]", " ", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r"-\n(?=[a-z])", "", text)  # de-hyphenate line breaks
    return "\n".join(line.strip() for line in text.split("\n")).strip()


def extract_text(filename: str, data: bytes) -> ExtractedDocument:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:  # pragma: no cover
            raise DocumentError("PDF support is not installed") from exc
        try:
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                try:
                    reader.decrypt("")
                except Exception as exc:  # noqa: BLE001
                    raise DocumentError("Encrypted PDFs are not supported") from exc
            pages = [(page.extract_text() or "") for page in reader.pages]
        except DocumentError:
            raise
        except Exception as exc:  # noqa: BLE001
            raise DocumentError("The PDF could not be parsed. Re-export it and try again.") from exc
        text = _clean("\n\n".join(pages))
        return ExtractedDocument(text=text, pages=len(pages))

    if suffix == ".docx":
        try:
            import docx
        except ImportError as exc:  # pragma: no cover
            raise DocumentError("DOCX support is not installed") from exc
        try:
            document = docx.Document(io.BytesIO(data))
        except Exception as exc:  # noqa: BLE001
            raise DocumentError("The DOCX file could not be parsed.") from exc
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text for cell in row.cells))
        text = _clean("\n".join(paragraphs))
        return ExtractedDocument(text=text, pages=max(1, len(paragraphs) // 40))

    if suffix in {".txt", ".md"}:
        try:
            text = data.decode("utf-8")
        except UnicodeDecodeError:
            text = data.decode("latin-1", errors="replace")
        return ExtractedDocument(text=_clean(text), pages=1)

    raise DocumentError("Only PDF, DOCX and TXT syllabus files are accepted")


def chunk_text(text: str, target_chars: int = 900, overlap: int = 150) -> list[str]:
    if not text:
        return []
    sentences = re.split(r"(?<=[.!?])\s+|\n{2,}", text)
    chunks: list[str] = []
    current = ""
    for sentence in sentences:
        sentence = sentence.strip()
        if not sentence:
            continue
        if len(current) + len(sentence) + 1 <= target_chars:
            current = f"{current} {sentence}".strip()
            continue
        if current:
            chunks.append(current)
            tail = current[-overlap:] if overlap else ""
            current = f"{tail} {sentence}".strip()
        else:
            # Very long single sentence: hard-split on word boundaries.
            words = sentence.split()
            buffer = ""
            for word in words:
                if len(buffer) + len(word) + 1 > target_chars:
                    chunks.append(buffer.strip())
                    buffer = word
                else:
                    buffer = f"{buffer} {word}".strip()
            current = buffer
    if current:
        chunks.append(current)
    return [chunk for chunk in chunks if len(chunk) >= 40]


def ingest_document(
    db: Session,
    *,
    subject: Subject,
    document: SyllabusDocument,
    data: bytes,
) -> int:
    extracted = extract_text(document.original_name, data)
    if len(extracted.text) < 80:
        raise DocumentError(
            "No usable text was extracted. Scanned PDFs need OCR before upload."
        )
    chunks = chunk_text(extracted.text)
    if not chunks:
        raise DocumentError("The document produced no usable text chunks.")
    embedder = get_embedder()
    vectors = embedder.embed(chunks)
    for index, (content, vector) in enumerate(zip(chunks, vectors)):
        db.add(
            DocumentChunk(
                document_id=document.id,
                subject_id=subject.id,
                chunk_index=index,
                content=content,
                char_count=len(content),
                embedding=vector,
                embedding_model=embedder.model,
            )
        )
    document.chunk_count = len(chunks)
    document.char_count = len(extracted.text)
    document.status = "PROCESSED"
    db.flush()
    return len(chunks)


def store_document_file(document: SyllabusDocument, filename: str, data: bytes) -> str:
    key = syllabus_key(document.subject_id, document.id, filename)
    return get_storage().put(key, data, document.content_type)


def retrieve(
    db: Session,
    *,
    subject_id: str,
    query: str,
    top_k: int = 6,
    min_score: float = 0.05,
) -> list[dict]:
    rows = (
        db.execute(
            select(DocumentChunk, SyllabusDocument.original_name)
            .join(SyllabusDocument, DocumentChunk.document_id == SyllabusDocument.id)
            .where(DocumentChunk.subject_id == subject_id)
        )
        .all()
    )
    if not rows:
        return []
    query_vector = get_embedder().embed_one(query)
    scored: list[dict] = []
    for chunk, document_name in rows:
        score = cosine(query_vector, chunk.embedding or [])
        if score >= min_score:
            scored.append(
                {
                    "chunk_id": chunk.id,
                    "document_id": chunk.document_id,
                    "document_name": document_name,
                    "content": chunk.content,
                    "score": round(score, 4),
                }
            )
    scored.sort(key=lambda item: item["score"], reverse=True)
    return scored[:top_k]


def document_fingerprint(document_id: str, filename: str) -> str:
    return hashlib.sha256(f"{document_id}:{filename}".encode()).hexdigest()[:12]


def subject_has_documents(db: Session, subject_id: str) -> bool:
    return (
        db.execute(
            select(SyllabusDocument.id).where(SyllabusDocument.subject_id == subject_id).limit(1)
        ).first()
        is not None
    )
